import json
import math
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


REPORT_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
REPORT_DISCLAIMER = (
    "AI-assisted report. Results support, but do not replace, recruiter judgement."
)
REPORT_FIELD_ORDER = (
    "email",
    "phone",
    "match_score",
    "component_scores",
    "summary",
    "matched_skills",
    "missing_skills",
    "gap_analysis",
    "recruitment_status",
)
REPORT_FIELDS = frozenset(REPORT_FIELD_ORDER)
MATCH_REPORT_FIELDS = (
    "candidate_id,status,recruitment_status,match_score,education_score,"
    "hard_skill_score,soft_skill_score,work_experience_score,summary,"
    "matched_skills,missing_skills,gap_analysis"
)
CANDIDATE_REPORT_FIELDS = "id,name,email,phone"


class ReportValidationError(ValueError):
    """Raised when requested report options are invalid."""


class ReportServiceError(RuntimeError):
    """A safe, user-facing report generation error."""


@dataclass(frozen=True)
class ReportArtifact:
    content: bytes
    filename: str
    ranked_candidate_count: int
    status_counts: dict[str, int]


def validate_report_fields(fields) -> tuple[str, ...]:
    if not isinstance(fields, list):
        raise ReportValidationError("Report fields must be an array.")
    if not all(isinstance(field, str) for field in fields):
        raise ReportValidationError("Each report field must be text.")

    requested = set(fields)
    invalid = requested - REPORT_FIELDS
    if invalid:
        raise ReportValidationError("One or more report fields are invalid.")
    return tuple(field for field in REPORT_FIELD_ORDER if field in requested)


def _response_data(response, default=None):
    data = getattr(response, "data", None)
    return default if data is None else data


def _valid_score(value) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        score = float(value)
    except (TypeError, ValueError):
        return None
    if not math.isfinite(score) or score < 0 or score > 100:
        return None
    return round(score, 2)


def load_job_report_data(
    supabase,
    recruiter_id: str,
    job_id: str,
) -> tuple[list[dict], dict[str, int]]:
    try:
        match_response = (
            supabase.table("match_results")
            .select(MATCH_REPORT_FIELDS)
            .eq("job_id", job_id)
            .execute()
        )
        matches = _response_data(match_response, [])
        completed_matches = []
        status_counts = {
            "pending": 0,
            "processing": 0,
            "completed": 0,
            "failed": 0,
        }

        for match in matches:
            status = str(match.get("status") or "").lower()
            if status in status_counts:
                status_counts[status] += 1
            score = _valid_score(match.get("match_score"))
            if status == "completed" and score is not None:
                completed_matches.append({**match, "match_score": score})

        candidate_ids = list(
            dict.fromkeys(
                str(match["candidate_id"])
                for match in completed_matches
                if match.get("candidate_id")
            )
        )
        candidates = {}
        if candidate_ids:
            candidate_response = (
                supabase.table("candidates")
                .select(CANDIDATE_REPORT_FIELDS)
                .eq("recruiter_id", recruiter_id)
                .in_("id", candidate_ids)
                .execute()
            )
            candidates = {
                str(candidate["id"]): candidate
                for candidate in _response_data(candidate_response, [])
            }
    except Exception as exc:
        raise ReportServiceError("Report data could not be loaded.") from exc

    ranked = []
    for match in completed_matches:
        candidate = candidates.get(str(match.get("candidate_id")))
        if not candidate:
            continue
        ranked.append({**candidate, **match})

    ranked.sort(
        key=lambda candidate: (
            -candidate["match_score"],
            str(candidate.get("name") or "").casefold(),
            str(candidate.get("candidate_id") or ""),
        )
    )
    return ranked, status_counts


def _readable_value(value) -> str:
    if value is None:
        return "--"
    if isinstance(value, list):
        cleaned = [str(item).strip() for item in value if str(item).strip()]
        return ", ".join(cleaned) if cleaned else "--"
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    text = str(value).strip()
    return text or "--"


def _display_recruitment_status(value) -> str:
    text = _readable_value(value)
    return text.replace("_", " ").title() if text != "--" else text


def _report_columns(fields: tuple[str, ...]):
    columns = [("Candidate Name", "name")]
    for field in fields:
        if field == "component_scores":
            columns.extend(
                [
                    ("Education Score", "education_score"),
                    ("Hard Skill Score", "hard_skill_score"),
                    ("Soft Skill Score", "soft_skill_score"),
                    ("Work Experience Score", "work_experience_score"),
                ]
            )
        else:
            labels = {
                "email": "Email",
                "phone": "Phone",
                "match_score": "Match Score",
                "summary": "Summary",
                "matched_skills": "Matched Skills",
                "missing_skills": "Missing Skills",
                "gap_analysis": "Gap Analysis",
                "recruitment_status": "Recruitment Status",
            }
            columns.append((labels[field], field))
    return columns


def _safe_filename(job_title: str, generated_at: datetime) -> str:
    safe_title = re.sub(r"[^A-Za-z0-9]+", "_", job_title).strip("_")
    safe_title = safe_title[:60] or "Job"
    return (
        f"TalentArch_{safe_title}_Candidate_Report_"
        f"{generated_at.strftime('%Y%m%d')}.xlsx"
    )


def build_job_report(
    job: dict,
    candidates: list[dict],
    status_counts: dict[str, int],
    fields: tuple[str, ...],
    *,
    generated_at: datetime | None = None,
) -> ReportArtifact:
    generated_at = generated_at or datetime.now(UTC)
    columns = _report_columns(fields)
    last_column = get_column_letter(max(2, len(columns)))

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Candidate Report"
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A8"

    blue = "1D5BF2"
    navy = "0F172A"
    slate = "475569"
    light_blue = "EFF6FF"
    light_slate = "F8FAFC"
    border_color = "E2E8F0"
    thin_border = Border(bottom=Side(style="thin", color=border_color))

    worksheet.merge_cells(f"A1:{last_column}1")
    title_cell = worksheet["A1"]
    title_cell.value = "TalentArch AI Candidate Report"
    title_cell.font = Font(name="Aptos Display", size=18, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor=blue)
    title_cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 34

    metadata = [
        ("Job Title", _readable_value(job.get("title"))),
        ("Generated", generated_at.astimezone(UTC).replace(tzinfo=None)),
        ("AI Disclaimer", REPORT_DISCLAIMER),
        (
            "Excluded from Ranking",
            (
                f"Pending: {status_counts.get('pending', 0)} | "
                f"Processing: {status_counts.get('processing', 0)} | "
                f"Failed: {status_counts.get('failed', 0)}"
            ),
        ),
    ]
    for row_index, (label, value) in enumerate(metadata, start=2):
        worksheet.cell(row=row_index, column=1, value=label).font = Font(
            name="Aptos", bold=True, color=navy
        )
        worksheet.merge_cells(
            start_row=row_index,
            start_column=2,
            end_row=row_index,
            end_column=max(2, len(columns)),
        )
        value_cell = worksheet.cell(row=row_index, column=2, value=value)
        value_cell.font = Font(name="Aptos", color=slate)
        value_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row_index == 3:
            value_cell.number_format = "yyyy-mm-dd hh:mm UTC"

    header_row = 7
    for column_index, (label, _) in enumerate(columns, start=1):
        cell = worksheet.cell(row=header_row, column=column_index, value=label)
        cell.font = Font(name="Aptos", bold=True, color=navy)
        cell.fill = PatternFill("solid", fgColor=light_blue)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin_border
    worksheet.row_dimensions[header_row].height = 28

    score_fields = {
        "match_score",
        "education_score",
        "hard_skill_score",
        "soft_skill_score",
        "work_experience_score",
    }
    wrapped_fields = {
        "summary",
        "matched_skills",
        "missing_skills",
        "gap_analysis",
    }

    if candidates:
        for row_index, candidate in enumerate(candidates, start=header_row + 1):
            for column_index, (_, field) in enumerate(columns, start=1):
                raw_value = candidate.get(field)
                if field in score_fields:
                    value = _valid_score(raw_value)
                elif field == "recruitment_status":
                    value = _display_recruitment_status(raw_value)
                else:
                    value = _readable_value(raw_value)

                cell = worksheet.cell(
                    row=row_index,
                    column=column_index,
                    value=value if value is not None else "--",
                )
                cell.font = Font(name="Aptos", color=slate)
                cell.alignment = Alignment(
                    wrap_text=field in wrapped_fields,
                    vertical="top",
                )
                cell.border = thin_border
                if field in score_fields and value is not None:
                    cell.number_format = '0.0"%"'
            if any(field in wrapped_fields for _, field in columns):
                worksheet.row_dimensions[row_index].height = 42
    else:
        empty_row = header_row + 1
        worksheet.merge_cells(
            start_row=empty_row,
            start_column=1,
            end_row=empty_row,
            end_column=max(1, len(columns)),
        )
        empty_cell = worksheet.cell(
            row=empty_row,
            column=1,
            value="No completed analyses available for ranking.",
        )
        empty_cell.fill = PatternFill("solid", fgColor=light_slate)
        empty_cell.font = Font(name="Aptos", italic=True, color=slate)
        empty_cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.row_dimensions[empty_row].height = 32

    data_end_row = max(header_row + 1, header_row + len(candidates))
    worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{data_end_row}"

    width_by_field = {
        "name": 24,
        "email": 30,
        "phone": 18,
        "match_score": 14,
        "education_score": 16,
        "hard_skill_score": 16,
        "soft_skill_score": 16,
        "work_experience_score": 22,
        "summary": 48,
        "matched_skills": 34,
        "missing_skills": 34,
        "gap_analysis": 48,
        "recruitment_status": 20,
    }
    for column_index, (_, field) in enumerate(columns, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = (
            width_by_field[field]
        )

    output = BytesIO()
    workbook.save(output)
    return ReportArtifact(
        content=output.getvalue(),
        filename=_safe_filename(str(job.get("title") or "Job"), generated_at),
        ranked_candidate_count=len(candidates),
        status_counts=status_counts,
    )


def generate_job_report(
    supabase,
    recruiter_id: str,
    job: dict,
    requested_fields,
    *,
    generated_at: datetime | None = None,
) -> ReportArtifact:
    fields = validate_report_fields(requested_fields)
    candidates, status_counts = load_job_report_data(
        supabase,
        recruiter_id,
        str(job["id"]),
    )
    return build_job_report(
        job,
        candidates,
        status_counts,
        fields,
        generated_at=generated_at,
    )
