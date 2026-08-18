import sys
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from docx import Document
from werkzeug.datastructures import FileStorage

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from services.resume_file_service import (  # noqa: E402
    ResumeFileError,
    extract_resume_text,
    validate_resume_file,
)

PDF_MIME = "application/pdf"
DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def make_file(content, filename, mime_type):
    return FileStorage(
        stream=BytesIO(content),
        filename=filename,
        content_type=mime_type,
    )


class FakePage:
    def __init__(self, text):
        self.text = text

    def extract_text(self):
        return self.text


class FakePdf:
    def __init__(self, pages):
        self.pages = pages

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return None


class ResumeFileServiceTests(unittest.TestCase):
    @patch("services.resume_file_service.pdfplumber.open")
    def test_valid_pdf_extracts_pages_in_order(self, open_pdf):
        open_pdf.return_value = FakePdf([FakePage("Page one"), FakePage("Page two")])
        resume = validate_resume_file(
            make_file(b"%PDF-test-content", "candidate.pdf", PDF_MIME)
        )

        self.assertEqual(extract_resume_text(resume), "Page one\n\nPage two")

    def test_valid_docx_extracts_paragraphs_and_tables(self):
        document = Document()
        document.add_paragraph("Candidate Name")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Python"
        content = BytesIO()
        document.save(content)

        resume = validate_resume_file(
            make_file(content.getvalue(), "candidate.docx", DOCX_MIME)
        )
        text = extract_resume_text(resume)

        self.assertIn("Candidate Name", text)
        self.assertIn("Skill | Python", text)

    def test_rejects_invalid_extension(self):
        with self.assertRaisesRegex(ResumeFileError, "Only PDF and DOCX"):
            validate_resume_file(make_file(b"text", "candidate.txt", "text/plain"))

    def test_rejects_mismatched_mime_type(self):
        with self.assertRaisesRegex(ResumeFileError, "MIME type"):
            validate_resume_file(make_file(b"%PDF-test", "candidate.pdf", "text/plain"))

    def test_rejects_empty_file(self):
        with self.assertRaisesRegex(ResumeFileError, "empty"):
            validate_resume_file(make_file(b"", "candidate.pdf", PDF_MIME))

    def test_rejects_oversized_file(self):
        with patch("services.resume_file_service.MAX_RESUME_FILE_SIZE", 5):
            with self.assertRaisesRegex(ResumeFileError, "20 MB"):
                validate_resume_file(make_file(b"%PDF-123", "candidate.pdf", PDF_MIME))

    def test_rejects_corrupted_pdf_signature(self):
        with self.assertRaisesRegex(ResumeFileError, "signature"):
            validate_resume_file(make_file(b"not a pdf", "candidate.pdf", PDF_MIME))

    def test_rejects_corrupted_docx(self):
        with self.assertRaisesRegex(ResumeFileError, "signature"):
            validate_resume_file(make_file(b"PK-not-a-zip", "candidate.docx", DOCX_MIME))

    @patch("services.resume_file_service.pdfplumber.open")
    def test_rejects_pdf_without_extractable_text(self, open_pdf):
        open_pdf.return_value = FakePdf([FakePage(""), FakePage(None)])
        resume = validate_resume_file(
            make_file(b"%PDF-test-content", "scan.pdf", PDF_MIME)
        )

        with self.assertRaisesRegex(ResumeFileError, "no extractable text"):
            extract_resume_text(resume)


if __name__ == "__main__":
    unittest.main()
