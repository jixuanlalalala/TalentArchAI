import sys
import unittest
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from services.report_service import (  # noqa: E402
    REPORT_DISCLAIMER,
    ReportValidationError,
    build_job_report,
    generate_job_report,
    validate_report_fields,
)


class FakeQuery:
    def __init__(self, response, operations):
        self.response = response
        self.operations = operations

    def select(self, fields):
        self.operations.append(("select", fields))
        return self

    def eq(self, field, value):
        self.operations.append(("eq", field, value))
        return self

    def in_(self, field, values):
        self.operations.append(("in", field, tuple(values)))
        return self

    def execute(self):
        return self.response


class FakeSupabase:
    def __init__(self, responses):
        self.responses = {name: list(items) for name, items in responses.items()}
        self.operations = []

    def table(self, name):
        self.operations.append(("table", name))
        return FakeQuery(self.responses[name].pop(0), self.operations)


def sample_matches():
    return [
        {
            "candidate_id": "candidate-1",
            "status": "completed",
            "recruitment_status": "shortlisted",
            "match_score": 84.5,
            "education_score": 80,
            "hard_skill_score": 90,
            "soft_skill_score": None,
            "work_experience_score": 75,
            "summary": "Strong backend alignment.",
            "matched_skills": ["Python", "PostgreSQL"],
            "missing_skills": ["Docker"],
            "gap_analysis": "Limited container experience.",
        },
        {
            "candidate_id": "candidate-2",
            "status": "completed",
            "recruitment_status": "new",
            "match_score": 92,
            "education_score": None,
            "hard_skill_score": 95,
            "soft_skill_score": 85,
            "work_experience_score": 90,
            "summary": "Best overall match.",
            "matched_skills": ["Python", "Docker"],
            "missing_skills": [],
            "gap_analysis": "No significant gaps identified.",
        },
        {"candidate_id": "candidate-3", "status": "pending", "match_score": None},
        {
            "candidate_id": "candidate-4",
            "status": "processing",
            "match_score": None,
        },
        {"candidate_id": "candidate-5", "status": "failed", "match_score": None},
        {
            "candidate_id": "candidate-6",
            "status": "completed",
            "match_score": None,
        },
    ]


class ReportServiceTests(unittest.TestCase):
    def test_report_fields_are_allowlisted_and_use_canonical_order(self):
        fields = validate_report_fields(
            ["gap_analysis", "email", "component_scores", "email"]
        )

        self.assertEqual(fields, ("email", "component_scores", "gap_analysis"))
        with self.assertRaises(ReportValidationError):
            validate_report_fields(["raw_text"])
        with self.assertRaises(ReportValidationError):
            validate_report_fields("email")

    def test_generation_ranks_only_completed_valid_owned_candidates(self):
        supabase = FakeSupabase(
            {
                "match_results": [SimpleNamespace(data=sample_matches())],
                "candidates": [
                    SimpleNamespace(
                        data=[
                            {
                                "id": "candidate-1",
                                "name": "Candidate One",
                                "email": "one@example.com",
                                "phone": "0111",
                            },
                            {
                                "id": "candidate-2",
                                "name": "Candidate Two",
                                "email": "two@example.com",
                                "phone": "0222",
                            },
                        ]
                    )
                ],
            }
        )
        generated_at = datetime(2026, 8, 10, 4, 0, tzinfo=UTC)

        report = generate_job_report(
            supabase,
            "recruiter-1",
            {"id": "job-1", "title": "Backend Engineer"},
            [
                "email",
                "match_score",
                "component_scores",
                "summary",
                "matched_skills",
                "missing_skills",
                "gap_analysis",
                "recruitment_status",
            ],
            generated_at=generated_at,
        )

        self.assertEqual(report.ranked_candidate_count, 2)
        self.assertEqual(
            report.status_counts,
            {"pending": 1, "processing": 1, "completed": 3, "failed": 1},
        )
        self.assertEqual(
            report.filename,
            "TalentArch_Backend_Engineer_Candidate_Report_20260810.xlsx",
        )
        workbook = load_workbook(BytesIO(report.content), data_only=True)
        self.assertEqual(workbook.sheetnames, ["Candidate Report"])
        sheet = workbook["Candidate Report"]
        headers = [cell.value for cell in sheet[7]]
        self.assertEqual(
            headers,
            [
                "Candidate Name",
                "Email",
                "Match Score",
                "Education Score",
                "Hard Skill Score",
                "Soft Skill Score",
                "Work Experience Score",
                "Summary",
                "Matched Skills",
                "Missing Skills",
                "Gap Analysis",
                "Recruitment Status",
            ],
        )
        self.assertEqual(sheet["A8"].value, "Candidate Two")
        self.assertEqual(sheet["C8"].value, 92)
        self.assertEqual(sheet["I8"].value, "Python, Docker")
        self.assertEqual(sheet["J8"].value, "--")
        self.assertEqual(sheet["L8"].value, "New")
        self.assertEqual(sheet["A9"].value, "Candidate One")
        self.assertEqual(sheet["D9"].value, 80)
        self.assertEqual(sheet["F9"].value, "--")
        self.assertEqual(sheet["B4"].value, REPORT_DISCLAIMER)
        self.assertIn("Pending: 1", sheet["B5"].value)
        self.assertIn(("eq", "recruiter_id", "recruiter-1"), supabase.operations)
        selected_fields = " ".join(
            operation[1]
            for operation in supabase.operations
            if operation[0] == "select"
        )
        for forbidden in ("raw_text", "resume_file_url", "location"):
            self.assertNotIn(forbidden, selected_fields)

    def test_empty_ranked_report_keeps_status_summary_and_message(self):
        report = build_job_report(
            {"id": "job-1", "title": "Engineer"},
            [],
            {"pending": 2, "processing": 0, "completed": 0, "failed": 1},
            ("match_score",),
            generated_at=datetime(2026, 8, 10, tzinfo=UTC),
        )

        workbook = load_workbook(BytesIO(report.content), data_only=True)
        sheet = workbook["Candidate Report"]
        self.assertEqual(
            sheet["A8"].value,
            "No completed analyses available for ranking.",
        )
        self.assertIn("Failed: 1", sheet["B5"].value)


if __name__ == "__main__":
    unittest.main()
